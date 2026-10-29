"""Quality-report rendering and export.

The report is stored as structured sections so the engineer can edit it inline. Exporters render
the same sections to HTML, PDF or Word. PDF/Word support is optional; the HTML export always works.
"""

from __future__ import annotations

import html
import io
import re
from collections.abc import Iterable
from datetime import UTC, datetime

from ..models import AuditEvent, Evidence, Report

BRAND = "RootCause"
AMBER = "#b07a25"


def _stamp() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def report_markdown(report: Report, *, evidence: Iterable[Evidence] = (), audits: Iterable[AuditEvent] = ()) -> str:
    lines = [
        f"# {report.title}",
        "",
        f"*Case {report.case_id} · report v{report.version} · {report.status.replace('_', ' ')}*",
        "",
        report.summary,
        "",
    ]
    for section in report.sections:
        lines += [f"## {section.heading}", "", section.body, ""]
    ev = list(evidence)
    if ev:
        lines += ["## Evidence appendix", ""]
        for e in ev:
            lines.append(f"- **{e.id} · {e.title}** — {e.detail} *({e.source})*")
        lines.append("")
    au = list(audits)
    if au:
        lines += ["## Audit trail", ""]
        for a in au:
            lines.append(f"- `{a.at}` **{a.actor}** {a.action} {a.object} — {a.detail}")
        lines.append("")
    lines += [
        "---",
        f"Prepared with {BRAND}. The engineer owns the approved cause; no automatic line control is performed.",
        "",
    ]
    return "\n".join(lines)


def _inline(text: str) -> str:
    text = html.escape(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"`(.+?)`", r"<code>\1</code>", text)
    return text.replace("\n", "<br>")


def report_html(
    report: Report,
    *,
    evidence: Iterable[Evidence] = (),
    audits: Iterable[AuditEvent] = (),
    printable: bool = True,
) -> str:
    sections = "".join(
        f'<section><h2>{html.escape(s.heading)}</h2><p>{_inline(s.body)}</p></section>' for s in report.sections
    )
    ev_html = "".join(
        f'<li><span class="eid">{html.escape(e.id)}</span><strong>{html.escape(e.title)}</strong>'
        f"<br>{_inline(e.detail)}<br><em>{html.escape(e.source)}</em></li>"
        for e in evidence
    )
    audit_html = "".join(
        f'<tr><td class="mono">{html.escape(a.at)}</td><td>{html.escape(a.actor)}</td>'
        f"<td>{html.escape(a.action)}</td><td>{html.escape(a.object)}</td><td>{_inline(a.detail)}</td></tr>"
        for a in audits
    )
    print_css = "@media print{body{background:#fff}.sheet{box-shadow:none;margin:0}.noprint{display:none}}" if printable else ""
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{html.escape(report.title)}</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
 :root{{color-scheme:light}}
 body{{margin:0;background:#e9e4d8;color:#22211d;font:15px/1.6 ui-sans-serif,system-ui,-apple-system,"Segoe UI",sans-serif}}
 .sheet{{max-width:820px;margin:28px auto;background:#f7f3ea;padding:46px 52px 36px;box-shadow:0 24px 70px rgba(0,0,0,.25);border-left:5px solid {AMBER};position:relative}}
 .kicker{{font:11px/1.4 ui-monospace,monospace;letter-spacing:.2em;text-transform:uppercase;color:#8a7f6c}}
 h1{{margin:8px 0 4px;font-size:31px;letter-spacing:-.02em}}
 .meta{{color:#6a6357;font-size:12px;margin-bottom:22px}}
 .summary{{font-size:16px;border-left:2px solid #d8cfbd;padding-left:14px;margin:20px 0 26px}}
 h2{{font-size:12px;text-transform:uppercase;letter-spacing:.14em;color:#8a7f6c;margin:26px 0 6px}}
 p{{margin:0 0 10px;white-space:normal}}
 .chain{{display:flex;gap:8px;flex-wrap:wrap;margin:14px 0 22px;font-size:12px}}
 .chain span{{border:1px solid #d8cfbd;background:#fffdf8;padding:6px 9px;border-radius:4px}}
 ul.evidence{{list-style:none;padding:0;margin:0}}
 ul.evidence li{{border-top:1px solid #e3dac8;padding:10px 0;font-size:13px}}
 .eid{{font:10px/1 ui-monospace,monospace;color:{AMBER};display:inline-block;margin-right:8px}}
 table{{width:100%;border-collapse:collapse;font-size:11px;margin-top:8px}}
 th,td{{text-align:left;border-bottom:1px solid #e3dac8;padding:6px 4px;vertical-align:top}}
 th{{color:#8a7f6c;font-weight:600;text-transform:uppercase;font-size:9px;letter-spacing:.1em}}
 .mono{{font-family:ui-monospace,monospace;font-size:10px}}
 .stamp{{position:absolute;right:46px;top:36px;border:2px solid #2f6b4f;color:#2f6b4f;padding:8px 12px;transform:rotate(-6deg);font:700 12px/1 ui-monospace,monospace;letter-spacing:.16em;opacity:.85}}
 footer{{margin-top:28px;border-top:1px solid #d8cfbd;padding-top:12px;font-size:10px;color:#8a7f6c}}
{print_css}
</style></head><body><main class="sheet">
 <div class="kicker">{html.escape(report.case_id)} · quality report · v{report.version}</div>
 <h1>{html.escape(report.title)}</h1>
 <div class="meta">Status: {html.escape(report.status.replace('_', ' '))} · Generated {html.escape(_stamp())}</div>
 {'<div class="stamp">APPROVED</div>' if report.status == 'approved' else ''}
 <div class="summary">{_inline(report.summary)}</div>
 {sections}
 {'<h2>Evidence appendix</h2><ul class="evidence">' + ev_html + '</ul>' if ev_html else ''}
 {'<h2>Audit trail</h2><table><tr><th>When</th><th>Who</th><th>Action</th><th>Object</th><th>Detail</th></tr>' + audit_html + '</table>' if audit_html else ''}
 <footer>Prepared with RootCause · evidence-first root cause analysis · the engineer owns the decision.
 No automatic line control is performed.</footer>
</main></body></html>"""


def report_pdf(report: Report, *, evidence: Iterable[Evidence] = ()) -> bytes | None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except Exception:
        return None
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(html.escape(report.title), styles["Title"]),
        Paragraph(f"{report.case_id} · v{report.version} · {report.status}", styles["Normal"]),
        Spacer(1, 8),
        Paragraph(html.escape(report.summary), styles["BodyText"]),
        Spacer(1, 10),
    ]
    for s in report.sections:
        story.append(Paragraph(html.escape(s.heading), styles["Heading2"]))
        story.append(Paragraph(html.escape(s.body).replace("\n", "<br/>"), styles["BodyText"]))
        story.append(Spacer(1, 6))
    ev = list(evidence)
    if ev:
        story.append(Paragraph("Evidence appendix", styles["Heading2"]))
        for e in ev:
            story.append(Paragraph(f"<b>{html.escape(e.id)} {html.escape(e.title)}</b> — {html.escape(e.detail)}", styles["BodyText"]))
    doc.build(story)
    return buf.getvalue()


def report_docx(report: Report, *, evidence: Iterable[Evidence] = ()) -> bytes | None:
    try:
        from docx import Document
    except Exception:
        return None
    doc = Document()
    doc.add_heading(report.title, level=1)
    doc.add_paragraph(f"{report.case_id} · report v{report.version} · {report.status}")
    doc.add_paragraph(report.summary)
    for s in report.sections:
        doc.add_heading(s.heading, level=2)
        for para in s.body.split("\n"):
            doc.add_paragraph(para)
    ev = list(evidence)
    if ev:
        doc.add_heading("Evidence appendix", level=2)
        for e in ev:
            doc.add_paragraph(f"{e.id} — {e.title}: {e.detail}", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def report_rtf(report: Report, *, evidence: Iterable[Evidence] = ()) -> bytes:
    """Dependency-free Word-compatible fallback."""

    def esc(t: str) -> str:
        return t.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}").replace("\n", "\\par ")

    parts = [r"{\rtf1\ansi\deff0", r"{\fonttbl{\f0 Segoe UI;}}", r"\fs36\b " + esc(report.title) + r"\b0\par", r"\fs20 " + esc(f"{report.case_id} · v{report.version} · {report.status}") + r"\par", esc(report.summary) + r"\par"]
    for s in report.sections:
        parts.append(r"\fs26\b " + esc(s.heading) + r"\b0\fs20\par " + esc(s.body) + r"\par")
    ev = list(evidence)
    if ev:
        parts.append(r"\fs26\b Evidence appendix\b0\fs20\par")
        for e in ev:
            parts.append(esc(f"• {e.id} {e.title} — {e.detail}") + r"\par")
    parts.append("}")
    return "\n".join(parts).encode("utf-8")
